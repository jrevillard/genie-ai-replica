# Copyright (C) 2026 International Telecommunication Union (ITU)
# SPDX-License-Identifier: Apache-2.0
"""xlsx format tests — the compatibility contract with tests/rag-benchmarks.

These pin the layouts VERBATIM: the suite format must survive a round trip
through OUR writer->reader AND be shaped exactly as xlsx_to_gold.py reads;
the results workbook must carry his NEW_COLS / COMPARE_HEADERS /
CALIBRATE_HEADERS names in his order.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

import xlsx_io


def _suite_fixture():
    return {
        "name": "Agri + Legal benchmark",
        "entries": [
            {
                "id": "new-1",
                "query": "What is the tomato planting season?",
                "reference_answer": "Tomatoes are planted after the last frost, typically early spring.",
                "reference_passages": [
                    "Plant tomatoes after the last frost date.",
                    "Season: early spring in temperate zones.",
                ],
                "source_doc": "agri-concepts/tomato-basics.md",
                "difficulty": "single_hop_specific",
                "language": "en",
                "oos": False,
            },
            {
                "id": "new-2",
                "query": "What is the capital of France?",
                "reference_answer": "",
                "reference_passages": [],
                "source_doc": "",
                "difficulty": "off_domain",
                "language": "en",
                "oos": True,
            },
        ],
    }


def test_suite_roundtrip_preserves_entries():
    suite = _suite_fixture()
    data = xlsx_io.write_suite_workbook(suite["name"], suite["entries"])
    parsed = xlsx_io.read_suite_workbook(data)
    entries = parsed["entries"]
    assert len(entries) == 2
    assert entries[0]["id"] == "new-1"
    assert entries[0]["query"] == "What is the tomato planting season?"
    assert entries[0]["reference_answer"].startswith("Tomatoes are planted")
    assert entries[0]["reference_passages"] == [
        "Plant tomatoes after the last frost date.",
        "Season: early spring in temperate zones.",
    ]
    assert entries[0]["source_doc"] == "agri-concepts/tomato-basics.md"
    assert entries[0]["oos"] is False
    assert entries[1]["oos"] is True
    assert entries[1]["reference_passages"] == []
    assert parsed["n_skipped_no_query"] == 0


def test_suite_layout_is_his_input_format():
    data = xlsx_io.write_suite_workbook("x", _suite_fixture()["entries"])
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb[xlsx_io.SUITE_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == tuple(xlsx_io.SUITE_HEADERS)  # header row 1, his names/order
    assert rows[1][0] == 1  # int id — his =ROW()-1 equivalence under data_only
    assert rows[1][2] == "What is the tomato planting season?"
    assert "Plant tomatoes after the last frost date.\n\nSeason:" in rows[1][4]  # blank-line separator
    assert rows[2][4] == "N/A"  # his OOS marker


def test_import_reads_his_hand_authored_sheet():
    """A sheet built the way his operators build them: sentinel-separated
    passages, an N/A out-of-scope row, a blank row, and a no-query row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Dataset"
    ws.append(["id", "difficulty", "query", "answer", "passages", "source"])
    ws.append([1, "easy", "Q one?", "A one.", "Passage one.@@@@@Passage two.", "doc-a"])
    ws.append([2, "", "Q two (out of scope)?", "", "N/A", ""])
    ws.append([None, None, None, None, None, None])  # blank row — skipped silently
    ws.append([4, "", "", "", "", ""])  # no query — skipped and counted
    ws.append([5, "", "Q five?", "", "Single passage", "doc-c"])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = xlsx_io.read_suite_workbook(buf.getvalue())
    entries = parsed["entries"]
    assert [e["id"] for e in entries] == ["new-1", "new-2", "new-5"]
    assert entries[0]["reference_passages"] == ["Passage one.", "Passage two."]
    assert entries[1]["oos"] is True
    assert entries[2]["reference_passages"] == ["Single passage"]
    assert parsed["n_skipped_no_query"] == 1


def test_import_accepts_non_test_dataset_sheet_via_active():
    wb = Workbook()
    ws = wb.active  # not named Test Dataset — active fallback
    ws.append(["id", "difficulty", "query", "answer", "passages", "source"])
    ws.append([1, "", "Only question", "Ref", "P", "s"])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = xlsx_io.read_suite_workbook(buf.getvalue())
    assert len(parsed["entries"]) == 1


def _run_fixture(extra_metric=None):
    result = {
        "id": "new-1",
        "answer": "Tomatoes are planted in early spring after the last frost.",
        "abstained": False,
        "context_keys": [f"chunk_{i}" for i in range(10)],
        "candidate_keys": ["c1", "c2"],
        "metrics": {
            "faithfulness": 0.92,
            "context_precision": 0.75,
            "context_recall": 0.5,
            "answer_relevancy": 0.88,
        },
    }
    if extra_metric:
        result["metrics"][extra_metric] = 0.64
    return {
        "label": "baseline",
        "params": {"judge": "granite"},
        "note": "first run",
        "results": [result],
        "summary": {"n_evaluable": 1, "n_abstained": 0, "n_unresolved": 0},
    }


def test_results_workbook_headers_are_his_verbatim():
    data = xlsx_io.write_results_workbook(_suite_fixture(), [_run_fixture()])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Gold", "Run_baseline", "Compare", "Calibrate"]
    run = wb["Run_baseline"]
    # row 1: config block; row 2: header; row 3: data (his 2-row preamble)
    assert run.cell(row=1, column=1).value == "params"
    assert run.cell(row=1, column=4).value == "note"
    header = [run.cell(row=2, column=c).value for c in range(1, len(xlsx_io.SUITE_HEADERS) + len(xlsx_io.NEW_COLS) + 1)]
    assert header == xlsx_io.SUITE_HEADERS + xlsx_io.NEW_COLS
    comp = [c.value for c in wb["Compare"][1]]
    assert comp == xlsx_io.COMPARE_HEADERS
    cal = [c.value for c in wb["Calibrate"][1]]
    assert cal == xlsx_io.CALIBRATE_HEADERS


def test_results_extended_metrics_append_after_his_22():
    data = xlsx_io.write_results_workbook(_suite_fixture(), [_run_fixture(extra_metric="answer_correctness")])
    wb = load_workbook(io.BytesIO(data))
    run = wb["Run_baseline"]
    ncols = len(xlsx_io.SUITE_HEADERS) + len(xlsx_io.NEW_COLS) + 1
    assert run.cell(row=2, column=ncols).value == "ragas_answer_correctness"
    assert run.cell(row=3, column=ncols).value == 0.64


def test_results_row_values_follow_his_conventions():
    data = xlsx_io.write_results_workbook(_suite_fixture(), [_run_fixture()])
    wb = load_workbook(io.BytesIO(data))
    run = wb["Run_baseline"]
    by_name = {run.cell(row=2, column=c).value: run.cell(row=3, column=c).value for c in range(1, 40)}
    assert by_name["ragas_faithfulness"] == 0.92
    assert by_name["gold_recall"] in (None, "")  # anchor columns blank (openpyxl reads "" back as None)
    assert by_name["abstained"] == "no"
    assert by_name["answer_first_200"].startswith("Tomatoes are planted")
    assert by_name["n_evaluable_helper"] == 1
    assert by_name["is_unresolved"] is False
    # keys cell: ';'-joined, 8 entries then ...(+N)
    assert by_name["selected_keys"] == ";".join(f"chunk_{i}" for i in range(8)) + "...(+2)"


def test_results_nan_metric_is_blank_cell():
    run = _run_fixture()
    run["results"][0]["metrics"]["faithfulness"] = None
    data = xlsx_io.write_results_workbook(_suite_fixture(), [run])
    wb = load_workbook(io.BytesIO(data))
    row = wb["Run_baseline"]
    by_name = {row.cell(row=2, column=c).value: row.cell(row=3, column=c).value for c in range(1, 40)}
    assert by_name["ragas_faithfulness"] in (None, "")  # blank cell either way


def test_export_neutralizes_formula_injection():
    # !622 round 21: openpyxl types '='-prefixed strings as FORMULAS — a
    # curator query starting '=' would execute when the workbook opens in
    # Excel/LibreOffice. The guard forces the string type with the content
    # preserved byte-for-byte. Discriminator: with data_only=True a stored
    # formula reads back as None (no cached value) — a forced string reads
    # back as its text.
    suite = _suite_fixture()
    suite["entries"][0]["query"] = '=HYPERLINK("http://evil.example","click me")'
    suite["entries"][0]["reference_answer"] = "=1+1"
    data = xlsx_io.write_suite_workbook(suite["name"], suite["entries"])
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb[xlsx_io.SUITE_SHEET]
    queries = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    answers = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
    assert queries[0] == '=HYPERLINK("http://evil.example","click me")'
    assert answers[0] == "=1+1"

    # The results workbook's Gold tab carries the same guard.
    run = _run_fixture()
    run["results"][0].setdefault("metrics", {})
    data2 = xlsx_io.write_results_workbook(suite, [run])
    wb2 = load_workbook(io.BytesIO(data2), data_only=True)
    gold = wb2["Gold"]
    assert gold.cell(row=2, column=3).value == '=HYPERLINK("http://evil.example","click me")'
