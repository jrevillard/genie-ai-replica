#!/usr/bin/env python3
# Copyright (C) 2025 ITU
# SPDX-License-Identifier: Apache-2.0
"""Convert a benchmark xlsx into a ``gold_dataset.json`` skeleton.

Generic tool — column mapping is configurable via CLI flags so the same script
handles any xlsx layout (El Salvador agri, Lesotho AI advisory, future domains).
The output schema matches what ``run_eval.py`` consumes (see ``gold_dataset.example.json``).

The first phase of the gold workflow:

    xlsx + schema   --(xlsx_to_gold.py)-->   gold_dataset.json (preview text only)
                                                    |
                          (after ArangoDB ingest)   v
                                          match_gold_chunks.py
                                                    |
                                                    v
                                          gold_dataset.json (content_hash + chunk_key filled)

content_hash / chunk_key are left empty here — they can only be filled AFTER the
corpus is ingested, because chunk _keys are auto-assigned at insert time.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    sys.stderr.write(
        "openpyxl is required. Install with: python3 -m venv /tmp/xlsx-venv && "
        "/tmp/xlsx-venv/bin/pip install openpyxl\n"
    )
    raise

# Crude EN/ES detector — Spanish typically has inverted punctuation marks
# (¿ ¡), ñ, or several accent-vowel tokens in the same short sentence.
_ES_TOKENS = re.compile(r"[¿¡ñ]|[áéíóú]", re.UNICODE)


def detect_language(text: str) -> str:
    """Heuristic: 2+ Spanish markers → 'es', otherwise 'en'."""
    if not text:
        return "en"
    hits = len(_ES_TOKENS.findall(text))
    return "es" if hits >= 2 else "en"


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--input-xlsx", required=True, type=Path, help="Source xlsx file")
    p.add_argument(
        "--output-json",
        type=Path,
        default=None,
        help="Output gold_dataset.json (default: <xlsx-stem>.gold.json next to input)",
    )
    p.add_argument("--sheet-name", default="Test Dataset", help="Worksheet name (default: %(default)s)")
    p.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="1-based row index of the column header (default: %(default)s)",
    )
    p.add_argument("--col-id", type=int, default=1, help="1-based column index for entry ID (default: 1)")
    p.add_argument("--col-query", type=int, default=3, help="1-based column index for the test query (default: 3)")
    p.add_argument(
        "--col-answer", type=int, default=4, help="1-based column index for the reference answer (default: 4)"
    )
    p.add_argument(
        "--col-passages", type=int, default=5,
        help="1-based column index for the verbatim relevant text passages (default: 5). "
             "Cells may contain multiple passages separated by a blank line — each becomes "
             "a separate expected_chunk.preview.",
    )
    p.add_argument(
        "--col-source", type=int, default=6,
        help="1-based column index for the document/section reference (default: 6). "
             "Optional but useful for traceability.",
    )
    p.add_argument(
        "--col-difficulty", type=int, default=2,
        help="1-based column index for the difficulty label (default: 2). Optional.",
    )
    p.add_argument(
        "--col-language", type=int,
        help="1-based column index for an explicit language column. If unset, auto-detect per query.",
    )
    p.add_argument(
        "--id-prefix", default="",
        help="String to prepend to every entry ID (e.g. 'q-') so multiple gold sets merge cleanly.",
    )
    p.add_argument(
        "--passage-separator", default=r"\n\s*\n",
        help="Regex that splits a passages cell into multiple previews (default: blank-line).",
    )
    p.add_argument(
        "--source-tag",
        action="append",
        default=[],
        help="Extra top-level key(s) to add under each entry (key=value). Repeatable. "
             "Use for 'language', 'domain', or any metadata that does not have its own xlsx column.",
    )
    return p.parse_args()


def split_passages(cell: str | None, separator: str) -> list[str]:
    """Split a 'Relevant Text Passages (Verbatim)' cell into clean preview strings.

    Returns an empty list for cells marked N/A / Out-of-Scope — the eval then
    scores these queries against an empty gold set (expected behaviour for
    unanswerable queries: the RAG pipeline MUST NOT surface any 'gold' chunk).
    """
    if not cell:
        return []
    text = cell.strip()
    if re.match(r"^(n/?a|out[-_ ]of[-_ ]scope|not\s+applicable)\b", text, re.IGNORECASE):
        return []
    parts = re.split(separator, text)
    return [p.strip() for p in parts if p.strip()]


def build_entry(
    row: tuple[Any, ...],
    args: argparse.Namespace,
    extra_tags: dict[str, str],
) -> dict[str, Any]:
    """Translate one xlsx row into a gold entry dict (no content_hash yet)."""
    raw_id = row[args.col_id - 1] if len(row) >= args.col_id else None
    query = row[args.col_query - 1] if len(row) >= args.col_query else None
    answer = row[args.col_answer - 1] if len(row) >= args.col_answer else None
    passages_cell = row[args.col_passages - 1] if len(row) >= args.col_passages else None
    source = row[args.col_source - 1] if len(row) >= args.col_source else None
    difficulty = row[args.col_difficulty - 1] if len(row) >= args.col_difficulty else None

    entry_id = f"{args.id_prefix}{raw_id}" if raw_id else f"{args.id_prefix}q?"
    lang_cell = row[args.col_language - 1] if args.col_language and len(row) >= args.col_language else None
    language = lang_cell.strip().lower()[:2] if lang_cell else detect_language(query or "")

    previews = split_passages(passages_cell, args.passage_separator)
    expected_chunks = [{"preview": p, "source_doc": (source or "").strip() or None} for p in previews]

    return {
        "id": entry_id,
        "query": (query or "").strip(),
        "language": language,
        "difficulty": (difficulty or "").strip() or None,
        "categoryLabels": [],  # domain-specific; populate via --source-tag or post-edit
        "serviceLabels": [],
        "reference_answer": (answer or "").strip(),
        "expected_chunks": expected_chunks,
        **extra_tags,
    }


def main() -> int:
    args = parse_args()
    if not args.input_xlsx.is_file():
        sys.stderr.write(f"ERROR: xlsx not found: {args.input_xlsx}\n")
        return 2

    extra_tags: dict[str, str] = {}
    for tag in args.source_tag:
        if "=" not in tag:
            sys.stderr.write(f"WARN: ignoring malformed --source-tag {tag!r} (want key=value)\n")
            continue
        k, v = tag.split("=", 1)
        extra_tags[k.strip()] = v.strip()

    wb = openpyxl.load_workbook(args.input_xlsx, data_only=True)
    if args.sheet_name not in wb.sheetnames:
        sys.stderr.write(
            f"ERROR: sheet {args.sheet_name!r} not in {args.input_xlsx.name}; "
            f"available: {wb.sheetnames}\n"
        )
        return 2
    ws = wb[args.sheet_name]

    rows_iter = ws.iter_rows(values_only=True)
    # Skip header
    for _ in range(args.header_row):
        next(rows_iter, None)

    entries: list[dict[str, Any]] = []
    skipped = 0
    for row in rows_iter:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        entry = build_entry(row, args, extra_tags)
        if not entry["query"]:
            skipped += 1
            continue
        entries.append(entry)

    out_path = args.output_json or args.input_xlsx.with_suffix(".gold.json")
    out_payload = {
        "_doc": (
            f"Auto-generated from {args.input_xlsx.name} (sheet {args.sheet_name!r}) by "
            "xlsx_to_gold.py. expected_chunks[].content_hash and chunk_key are filled by "
            "match_gold_chunks.py AFTER the corpus is ingested into ArangoDB. Schema matches "
            "gold_dataset.example.json."
        ),
        "source_xlsx": str(args.input_xlsx.resolve()),
        "source_sheet": args.sheet_name,
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "n_entries": len(entries),
        "n_skipped_no_query": skipped,
        "entries": entries,
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(out_payload, indent=2, ensure_ascii=False))
    print(
        f"Wrote {len(entries)} entries ({skipped} skipped for missing query) to {out_path}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
